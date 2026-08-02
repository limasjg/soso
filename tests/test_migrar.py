from datetime import date

import openpyxl
from sqlalchemy import select

from decimal import Decimal

from database import Categoria, Lancamento
from migrar import extrair_lancamentos, importar, para_decimal


def criar_planilha(caminho):
    wb = openpyxl.Workbook()
    aba = wb.active
    aba.title = "2026"
    # Primeiro bloco mensal: E/F. As linhas de totais devem ser ignoradas.
    aba.cell(4, 5, "Aluguel")
    aba.cell(4, 6, 2500)
    aba.cell(5, 5, "Ganho")
    aba.cell(5, 6, 7000)
    aba.cell(6, 5, "Investido")
    aba.cell(6, 6, 1000)
    aba.cell(7, 5, "Contas")
    aba.cell(7, 6, 2500)
    # Segundo bloco mensal: I/J.
    aba.cell(4, 9, "Luz")
    aba.cell(4, 10, "R$ 120,50")
    wb.save(caminho)


def test_para_decimal_aceita_formato_brasileiro():
    assert para_decimal("R$ 2.504,08") == Decimal("2504.08")
    assert para_decimal("inválido") is None


def test_importacao_e_idempotente(tmp_path):
    arquivo = tmp_path / "entrada.xlsx"
    criar_planilha(arquivo)
    registros = extrair_lancamentos(arquivo)
    assert len(registros) == 4
    url = f"sqlite:///{tmp_path / 'teste.db'}"
    assert importar(arquivo, url) == 4
    assert importar(arquivo, url) == 0

    from database import criar_engine, criar_sessao
    Sessao = criar_sessao(criar_engine(url))
    with Sessao() as sessao:
        lancamentos = sessao.scalars(select(Lancamento).order_by(Lancamento.data)).all()
        assert {item.tipo for item in lancamentos} == {"receita", "despesa", "investimento"}
        assert lancamentos[0].data == date(2026, 1, 1)
        assert sessao.scalars(select(Categoria)).all()


def test_aba_anual_vazia_nao_interrompe_importacao(tmp_path):
    arquivo = tmp_path / "vazia.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "2027"
    wb.save(arquivo)
    assert extrair_lancamentos(arquivo) == []
