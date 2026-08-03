"""Importa a planilha histórica para o banco de dados do SOSO.

Uso: python migrar.py --arquivo financas.xlsx
O script pode ser executado mais de uma vez: registros previamente importados
não são inseridos novamente.
"""
from __future__ import annotations

import argparse
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from sqlalchemy import delete, select

from database import Categoria, GastoFixo, Lancamento, criar_engine, criar_sessao, criar_tabelas

MESES = {nome: numero for numero, nome in enumerate(
    ("janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"), 1
)}
MESES["outurbro"] = 10  # grafia presente na aba de 2022
IGNORAR = {"pessoal", "contas", "total", "valor", "mês", "ganho total", "valor total"}
ROTULOS_RECEITA = {"ganho", "receita", "salário", "salario"}
ROTULOS_INVESTIMENTO = {"investido", "investimento", "investimentos"}


def texto(valor: object) -> str:
    return str(valor or "").strip()


def normalizar(valor: object) -> str:
    return re.sub(r"\s+", " ", texto(valor).lower())


def para_decimal(valor: object) -> Decimal | None:
    """Converte números e textos como 'R$ 2.504,08' sem aceitar lixo."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    bruto = texto(valor).replace("R$", "").replace(" ", "")
    if not bruto:
        return None
    if "," in bruto:
        bruto = bruto.replace(".", "").replace(",", ".")
    try:
        return Decimal(bruto).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def tipo_do_rotulo(rotulo: str) -> str:
    chave = normalizar(rotulo)
    if chave in ROTULOS_RECEITA:
        return "receita"
    if chave in ROTULOS_INVESTIMENTO:
        return "investimento"
    return "despesa"


def abas_anuais(workbook) -> list[tuple[int, object]]:
    resultado = []
    for aba in workbook.worksheets:
        try:
            ano = int(aba.title)
        except ValueError:
            continue
        if 2000 <= ano <= 2100:
            resultado.append((ano, aba))
    return resultado


def cabecalhos_mensais(aba) -> list[tuple[int, int, int]]:
    """Localiza os blocos cujo cabeçalho é um mês e cuja tabela começa abaixo."""
    encontrados = []
    for linha in aba.iter_rows():
        for celula in linha:
            mes = MESES.get(normalizar(celula.value))
            if not mes:
                continue
            if (
                normalizar(aba.cell(celula.row + 1, celula.column).value) == "contas"
                and normalizar(aba.cell(celula.row + 1, celula.column + 1).value) == "valor"
            ):
                encontrados.append((mes, celula.row, celula.column))
    return encontrados


def extrair_bloco_mensal(aba, ano: int, mes: int, linha_cabecalho: int, coluna: int) -> list[tuple]:
    """Extrai contas e o resumo de ganho/investimento de um bloco mensal."""
    registros = []
    linha = linha_cabecalho + 2

    # A tabela de contas é contínua e termina na primeira linha vazia.
    while linha <= aba.max_row:
        descricao = texto(aba.cell(linha, coluna).value)
        valor = para_decimal(aba.cell(linha, coluna + 1).value)
        if not descricao and valor is None:
            break
        chave = normalizar(descricao)
        if descricao and valor is not None and valor > 0 and chave not in IGNORAR:
            registros.append((ano, mes, descricao, valor, "despesa"))
        linha += 1

    # Depois da linha vazia há o resumo. Nele, somente Ganho e Investido são
    # lançamentos; Pessoal e Contas são totais e não devem ser duplicados.
    for linha_resumo in range(linha + 1, (aba.max_row or 0) + 1):
        descricao = texto(aba.cell(linha_resumo, coluna).value)
        chave = normalizar(descricao)
        if chave in MESES:
            break
        if chave not in ROTULOS_RECEITA | ROTULOS_INVESTIMENTO:
            continue
        valor = para_decimal(aba.cell(linha_resumo, coluna + 1).value)
        if valor is not None and valor > 0:
            registros.append((ano, mes, descricao, valor, tipo_do_rotulo(descricao)))
    return registros


def extrair_formato_legado(aba, ano: int) -> list[tuple]:
    """Mantém suporte às planilhas antigas sem cabeçalhos de mês."""
    registros = []
    for mes in range(1, 13):
        coluna_descricao = 5 + (mes - 1) * 4
        for linha in range(4, (aba.max_row or 0) + 1):
            descricao = texto(aba.cell(linha, coluna_descricao).value)
            valor = para_decimal(aba.cell(linha, coluna_descricao + 1).value)
            chave = normalizar(descricao)
            if not descricao or valor is None or valor <= 0 or chave in IGNORAR or chave in MESES:
                continue
            registros.append((ano, mes, descricao, valor, tipo_do_rotulo(descricao)))
    return registros


def extrair_lancamentos(arquivo: str | Path):
    """Lê os blocos mensais da planilha e devolve tuplas prontas para importar."""
    # A planilha original não declara a dimensão usada em todas as abas; por
    # isso, o modo read_only não consegue informar max_row com confiabilidade.
    wb = openpyxl.load_workbook(arquivo, data_only=True, read_only=False)
    registros = []
    for ano, aba in abas_anuais(wb):
        cabecalhos = cabecalhos_mensais(aba)
        if cabecalhos:
            for mes, linha, coluna in cabecalhos:
                registros.extend(extrair_bloco_mensal(aba, ano, mes, linha, coluna))
        else:
            registros.extend(extrair_formato_legado(aba, ano))
    wb.close()
    return registros


def obter_categoria(sessao, nome: str, tipo: str) -> Categoria:
    categoria = sessao.scalar(select(Categoria).where(Categoria.nome == nome))
    if categoria is None:
        categoria = Categoria(nome=nome, tipo=tipo)
        sessao.add(categoria)
        sessao.flush()
    return categoria


def importar(arquivo: str | Path, url: str | None = None, substituir_ano: int | None = None) -> int:
    engine = criar_engine(url)
    criar_tabelas(engine)
    Sessao = criar_sessao(engine)
    inseridos = 0
    with Sessao.begin() as sessao:
        registros = extrair_lancamentos(arquivo)
        if substituir_ano is not None:
            registros = [registro for registro in registros if registro[0] == substituir_ano]
        if substituir_ano is not None:
            sessao.execute(delete(Lancamento).where(
                Lancamento.origem == "excel",
                Lancamento.data >= date(substituir_ano, 1, 1),
                Lancamento.data < date(substituir_ano + 1, 1, 1),
            ))
        for ano, mes, descricao, valor, tipo in registros:
            categoria = obter_categoria(sessao, descricao, tipo)
            data_lancamento = date(ano, mes, 1)
            existente = sessao.scalar(select(Lancamento.id).where(
                Lancamento.descricao == descricao,
                Lancamento.tipo == tipo,
                Lancamento.valor == valor,
                Lancamento.data == data_lancamento,
                Lancamento.origem == "excel",
            ))
            if existente:
                continue
            gasto_fixo_id = None
            if tipo == "despesa":
                fixo = sessao.scalar(select(GastoFixo).where(GastoFixo.descricao == descricao))
                if fixo is None:
                    fixo = GastoFixo(categoria_id=categoria.id, descricao=descricao, valor_previsto=valor, dia_vencimento=10)
                    sessao.add(fixo)
                    sessao.flush()
                gasto_fixo_id = fixo.id
            sessao.add(Lancamento(categoria_id=categoria.id, tipo=tipo, descricao=descricao,
                                  valor=valor, data=data_lancamento, gasto_fixo_id=gasto_fixo_id,
                                  origem="excel"))
            inseridos += 1
    return inseridos


def main() -> None:
    parser = argparse.ArgumentParser(description="Importa financas.xlsx para o Supabase.")
    parser.add_argument("--arquivo", default="financas.xlsx", help="Caminho da planilha de origem")
    parser.add_argument("--substituir-ano", type=int, help="Remove os lançamentos de origem Excel do ano antes de importá-lo novamente")
    args = parser.parse_args()
    if not Path(args.arquivo).exists():
        parser.error(f"Arquivo não encontrado: {args.arquivo}")
    total = importar(args.arquivo, substituir_ano=args.substituir_ano)
    print(f"Importação concluída: {total} lançamento(s) novos.")


if __name__ == "__main__":
    main()
